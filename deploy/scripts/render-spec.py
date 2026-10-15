#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""이관 명세서(deploy/SPEC.md)를 Word(.docx)·PDF 로 뽑는다.

    python3 deploy/scripts/render-spec.py OUTDIR            # docx + pdf
    python3 deploy/scripts/render-spec.py OUTDIR --no-pdf   # docx 만

원본은 SPEC.md 하나뿐이다. **Word·PDF 를 손으로 고치지 말 것** — SPEC.md 를 고치고 다시 뽑는다.

두 파일은 같은 파서(parse)에서 나오지만 그리는 엔진이 다르다.
  · .docx — python-docx 로 직접 쓴다. 글꼴 맑은 고딕 + Consolas(Windows 기본, Mac 용 Word 에도 동봉).
           목차는 TOC 필드 안에 항목·내부 링크를 미리 채워 둔다(쪽번호 없음). Word 에서 목차를 선택하고
           F9 를 누르면 Word 가 쪽번호까지 다시 만든다.
  · .pdf  — 같은 블록을 HTML 로 옮겨 Chrome(또는 Edge) 헤드리스로 인쇄한다. 글꼴은 레포의
           tools/NotoSansKR.ttf 를 박아 넣으므로 어느 PC 에서 뽑아도 같다.
           목차 쪽번호는 한 번 인쇄한 PDF 에서 각 제목의 쪽을 찾아 넣고 다시 인쇄한다 — 쪽이 바뀌지
           않을 때까지(pypdf).

왜 Word 로 PDF 를 만들지 않나
  이 맥의 Word 는 라이선스 미활성(보기 전용)이라 [다른 이름으로 저장]이 막혀 있고(AppleScript save as → -1708),
  Mac Word AppleScript 는 TOC 필드 갱신도 하지 못한다(update → -1708, update field → false). 2026-10-15 실측.

필요한 것: python-docx · pypdf (pip install python-docx pypdf) · Chrome 또는 Edge (경로는 SL_CHROME 로 지정 가능)
"""
import argparse, html, os, pathlib, re, shutil, subprocess, sys, tempfile, unicodedata

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor, Twips
except ImportError:
    sys.exit("✗ python-docx 가 없다 — pip install python-docx")

HERE = os.path.dirname(os.path.abspath(__file__))
SPEC = os.path.join(os.path.dirname(HERE), "SPEC.md")
BASENAME = "쉴더스랩_일반서버_이관명세서"

FONT, MONO = "Malgun Gothic", "Consolas"
PINE, INK, MUTED = RGBColor(0x1A, 0x4B, 0x3A), RGBColor(0x1C, 0x1C, 0x1A), RGBColor(0x6B, 0x6A, 0x63)
PAGE_W, MARGIN_LR = Mm(210), Mm(18)
TEXT_W = Mm(210 - 2 * 18)                          # 본문 폭(표 폭 계산 기준) — Length 끼리 빼면 int 가 된다


# ── 마크다운 → 블록 ─────────────────────────────────────────────────────────
# SPEC.md 가 실제로 쓰는 문법만 다룬다: 제목 · 문단 · 표 · 코드 블록 · 인용 · 목록 · 구분선.
# 글자 누락은 verify() 가 원문 조각 단위로 잡는다. 모양(줄바꿈·쪽 나눔·정렬)은 페이지 이미지로 확인한다.
MARK = ("⚠", "🚨", "🔴", "🟡", "✅", "⏳", "⛔", "- ", "· ")

def cat(a, b):
    """원문 줄바꿈을 이어 붙인다 — 마크다운은 공백으로 잇지만 가운뎃점 나열(A·B·C)은 붙여 쓰므로
    줄이 '·' 로 끝나거나 시작하면 공백을 넣지 않는다("길이 상한· IP" 처럼 보이지 않게)."""
    return a + ("" if a.endswith("·") or b.startswith("·") else " ") + b

def _join(lines):
    """줄바꿈은 공백으로 잇되(마크다운 규칙), 표식으로 시작하는 줄은 새 줄로 둔다."""
    out = []
    for ln in lines:
        if out and not ln.startswith(MARK):
            out[-1] = cat(out[-1], ln)
        else:
            out.append(ln)
    return out

def _cells(row):
    row = row.strip()
    if row.startswith("|"): row = row[1:]
    if row.endswith("|") and not row.endswith("\\|"): row = row[:-1]
    cells, cur, code, i = [], "", False, 0
    while i < len(row):
        ch = row[i]
        if ch == "\\" and i + 1 < len(row) and row[i + 1] == "|":
            cur += "\\|"; i += 2; continue
        if ch == "`": code = not code
        if ch == "|" and not code:
            cells.append(cur.strip()); cur = ""
        else:
            cur += ch
        i += 1
    cells.append(cur.strip())
    return cells

def parse(md):
    lines = md.splitlines()
    blocks, i, n = [], 0, len(lines)
    while i < n:
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        if ln.startswith("```"):
            body, i = [], i + 1
            while i < n and not lines[i].startswith("```"):
                body.append(lines[i]); i += 1
            blocks.append(("code", body)); i += 1; continue
        m = re.match(r"^(#{1,6})\s+(.*)$", ln)
        if m:
            blocks.append(("h", len(m.group(1)), m.group(2).strip())); i += 1; continue
        if re.match(r"^-{3,}\s*$", ln):
            blocks.append(("hr",)); i += 1; continue
        if ln.lstrip().startswith("|"):
            rows = []
            while i < n and lines[i].lstrip().startswith("|"):
                rows.append(lines[i]); i += 1
            cells = [_cells(r) for r in rows]
            if len(cells) > 1 and all(re.match(r"^:?-+:?$", c) for c in cells[1] if c):
                cells.pop(1)
            blocks.append(("table", cells)); continue
        if ln.startswith(">"):
            paras, cur = [], []
            while i < n and lines[i].startswith(">"):
                t = lines[i][1:].strip()
                if t: cur.append(t)
                elif cur: paras.append(cur); cur = []
                i += 1
            if cur: paras.append(cur)
            blocks.append(("quote", [_join(p) for p in paras])); continue
        lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", ln)
        if lm:
            items = []
            while i < n:
                lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", lines[i])
                if lm:
                    text = re.sub(r"^\[ \]\s+", "□ ", lm.group(3))     # 체크리스트 항목
                    items.append([lm.group(2), text]); i += 1
                elif lines[i].startswith("  ") and lines[i].strip() and items:
                    items[-1][1] = cat(items[-1][1], lines[i].strip()); i += 1
                else:
                    break
            blocks.append(("list", items)); continue
        para = []
        while i < n and lines[i].strip() and not re.match(r"^(#{1,6}\s|```|>|\s*\||-{3,}\s*$|\s*([-*]|\d+\.)\s)", lines[i]):
            para.append(lines[i].strip()); i += 1
        text = para[0]
        for ln in para[1:]: text = cat(text, ln)
        blocks.append(("p", text))
    return blocks


# ── 인라인: **굵게** · *기울임* · `코드` · \| ────────────────────────────────
def inline(text):
    segs, buf, bold, ital, i, n = [], "", False, False, 0, len(text)
    def flush():
        nonlocal buf
        if buf: segs.append((buf, bold, ital, False)); buf = ""
    while i < n:
        c = text[i]
        if c == "\\" and i + 1 < n and text[i + 1] in "|*`\\_":
            buf += text[i + 1]; i += 2; continue
        if c == "`":
            j = text.find("`", i + 1)
            if j > i:
                flush(); segs.append((text[i + 1:j].replace("\\|", "|"), bold, ital, True)); i = j + 1; continue
        if text.startswith("**", i):
            flush(); bold = not bold; i += 2; continue
        if c == "*":
            opening = not ital and i + 1 < n and not text[i + 1].isspace() and text.find("*", i + 1) > 0
            closing = ital and i > 0 and not text[i - 1].isspace()
            if opening or closing:
                flush(); ital = not ital; i += 1; continue
        buf += c; i += 1
    flush()
    return segs

def plain(text):
    return "".join(s[0] for s in inline(text))


# ── docx 저수준 도우미 ───────────────────────────────────────────────────────
def set_fonts(rpr_owner, latin=FONT, ea=FONT):
    """테마 글꼴 속성을 지우고 라틴·한글 글꼴을 명시한다(안 지우면 테마가 이긴다)."""
    rpr = rpr_owner.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    rf.set(qn("w:ascii"), latin); rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), ea); rf.set(qn("w:cs"), latin)

def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items(): e.set(qn("w:" + k), str(v))
    return e

# OOXML 은 속성 요소의 **순서**를 스키마로 정한다(xsd:sequence). Word 는 순서가 틀려도 열지만
# 엄격한 처리기는 "읽을 수 없는 내용"으로 본다 → 스키마 순서대로 끼워 넣는다(ISO 29500 wml.xsd).
ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
           "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection "
           "textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
             "tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription tblPrChange",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark "
            "headers cellIns cellDel cellMerge tcPrChange",
}

def put(parent, child):
    """parent(pPr·tblPr·tcPr) 안에 child 를 스키마 순서 자리에 넣는다. 같은 요소가 있으면 바꾼다."""
    seq = ORDER[parent.tag.split("}")[1]].split()
    name = child.tag.split("}")[1]
    old = parent.find(qn("w:" + name))
    if old is not None: parent.remove(old)
    after = seq[seq.index(name) + 1:]
    for k, ex in enumerate(parent):
        if ex.tag.split("}")[1] in after:
            parent.insert(k, child); return child
    parent.append(child); return child

def shade(ppr_or_tcpr, fill):
    put(ppr_or_tcpr, el("w:shd", val="clear", color="auto", fill=fill))

def para_border(p, sides, color, sz=6, space=4):
    b = el("w:pBdr")
    for s in sides: b.append(el("w:" + s, val="single", sz=sz, space=space, color=color))
    put(p._p.get_or_add_pPr(), b)

def add_runs(p, text, size=None, mono_size=None, color=None, force_bold=False):
    for t, b, it, code in inline(text):
        r = p.add_run(t)
        r.bold = b or force_bold or None
        r.italic = it or None
        if code:
            set_fonts(r._r, MONO, FONT)
            r.font.size = Pt(mono_size or (size - 0.5 if size else 9))
            r.font.color.rgb = RGBColor(0x2E, 0x3B, 0x35) if color is None else color
        else:
            if size: r.font.size = Pt(size)
            if color is not None: r.font.color.rgb = color

def field(p, instr, placeholder=""):
    r = p.add_run(); r._r.append(el("w:fldChar", fldCharType="begin"))   # dirty 를 붙이면 Word 가 열 때 "필드 업데이트?" 창을 띄운다
    r = p.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._r.append(it)
    r = p.add_run(); r._r.append(el("w:fldChar", fldCharType="separate"))
    p.add_run(placeholder)
    r = p.add_run(); r._r.append(el("w:fldChar", fldCharType="end"))


# ── 표 폭: 내용 길이로 비례 배분 ─────────────────────────────────────────────
def dispw(s):
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in s)

def col_widths(rows, total):
    ncol = max(len(r) for r in rows)
    U, PAD = 82, 200                                 # 9pt 기준 반각 1자 ≈ 82 twips · 셀 여백
    minw, want = [], []
    for c in range(ncol):
        texts = [plain(r[c]) if c < len(r) else "" for r in rows]
        longest_tok = max((dispw(t) for s in texts for t in s.split()), default=1)
        longest = max((dispw(s) for s in texts), default=1)
        mn = min(longest if longest <= 12 else longest_tok, 26) * U + PAD   # 짧은 칸(≤12칸)은 한 줄로
        minw.append(max(mn, 520))
        want.append(max(minw[-1], min(longest, 64) * U + PAD))
    if sum(want) <= total:
        k = total / sum(want); return [int(w * k) for w in want]
    extra = total - sum(minw)
    if extra <= 0:
        k = total / sum(minw); return [int(w * k) for w in minw]
    span = sum(w - m for w, m in zip(want, minw)) or 1
    return [int(m + extra * (w - m) / span) for w, m in zip(want, minw)]


# ── 문서 ────────────────────────────────────────────────────────────────────
def setup(doc):
    # python-docx 기본 서식은 Word 2010 호환 모드로 열린다 → 현재 모드(15)로 표시
    settings = doc.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = el("w:compat"); settings.append(compat)
    for cs in compat.findall(qn("w:compatSetting")):
        if cs.get(qn("w:name")) == "compatibilityMode": compat.remove(cs)
    compat.append(el("w:compatSetting", name="compatibilityMode", uri="http://schemas.microsoft.com/office/word", val=15))
    zoom = settings.find(qn("w:zoom"))                  # 기본 서식의 zoom 에 필수 속성 percent 가 빠져 있다
    if zoom is not None and zoom.get(qn("w:percent")) is None: zoom.set(qn("w:percent"), "100")
    s = doc.sections[0]
    s.page_width, s.page_height = PAGE_W, Mm(297)
    s.left_margin = s.right_margin = MARGIN_LR
    s.top_margin, s.bottom_margin = Mm(20), Mm(18)
    s.header_distance = s.footer_distance = Mm(9)
    st = doc.styles
    normal = st["Normal"]
    set_fonts(normal.element); normal.font.size = Pt(10); normal.font.color.rgb = INK
    rpr = normal.element.get_or_add_rPr(); rpr.append(el("w:lang", val="en-US", eastAsia="ko-KR"))
    pf = normal.paragraph_format; pf.space_after = Pt(4); pf.line_spacing = 1.3
    for name, size, before, after, color in (("Heading 1", 16, 20, 8, PINE), ("Heading 2", 12.5, 14, 5, INK),
                                             ("Heading 3", 11, 10, 4, INK), ("Heading 4", 10.5, 8, 3, INK)):
        h = st[name]; set_fonts(h.element)
        h.font.size, h.font.bold, h.font.italic, h.font.color.rgb = Pt(size), True, False, color
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(before), Pt(after)
        h.paragraph_format.keep_with_next = True; h.paragraph_format.line_spacing = 1.15
    for name in ("Title", "List Bullet", "List Number"):
        set_fonts(st[name].element)
    for name in ("Header", "Footer"):
        try: set_fonts(st[name].element)
        except KeyError: pass

def header_footer(doc, label):
    s = doc.sections[0]; s.different_first_page_header_footer = True
    hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run(label); r.font.size = Pt(8); r.font.color.rgb = MUTED
    para_border(hp, ["bottom"], "D9D6CC", sz=4, space=4)
    fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(fp, "PAGE", "1"); fp.add_run(" / "); field(fp, "NUMPAGES", "1")
    for rr in fp.runs: rr.font.size = Pt(8); rr.font.color.rgb = MUTED

def add_table(doc, rows, header=True, size=8.5, widths=None):
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    put(tblPr, el("w:tblLayout", type="fixed"))
    b = el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(el("w:" + side, val="single", sz=4, space=0, color="CFCBC0"))
    put(tblPr, b)
    mar = el("w:tblCellMar")
    for side, w in (("top", 45), ("left", 90), ("bottom", 45), ("right", 90)):   # CT_TblCellMar 순서
        mar.append(el("w:" + side, w=w, type="dxa"))
    put(tblPr, mar)
    widths = widths or col_widths(rows, int(TEXT_W.twips))
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(w))
    for ri, row in enumerate(rows):
        tr = t.rows[ri]._tr
        trPr = tr.get_or_add_trPr()
        if ri == 0 and header: trPr.append(el("w:tblHeader"))
        trPr.append(el("w:cantSplit"))
        for ci, text in enumerate(row):
            cell = t.cell(ri, ci); cell.width = Twips(widths[ci])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.2
            if ri == 0 and header:
                shade(cell._tc.get_or_add_tcPr(), "1A4B3A")
                add_runs(p, text, size=size, color=RGBColor(0xFF, 0xFF, 0xFF), force_bold=True)
            else:
                add_runs(p, text, size=size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def add_code(doc, lines):
    """코드 한 줄 = 문단 하나. 테두리·음영이 같은 문단은 Word 가 한 상자로 묶는다.
    긴 줄이 접히면 이어지는 줄을 HANG 칸 들여 써서(걸린 들여쓰기) 새 명령처럼 보이지 않게 한다."""
    size = 7.8
    # 도식 정렬: 원문은 "한글 1자 = 영문 2칸" 으로 맞춰져 있다. Consolas 는 0.5498em, 맑은 고딕 한글은 1.0em 이라
    # 그대로 두면 한글이 든 줄마다 0.1em 씩 밀린다 → 한글(전각)에만 자간을 더해 정확히 2칸으로 맞춘다.
    # (상자·화살표 문자는 Consolas 에 있어 이미 1칸이다 — 2026-10-15 fontTools 로 실측)
    wide_pad = round((2 * CONSOLAS_ADV - MALGUN_HANGUL_ADV) * size * 20) if is_diagram(lines) else 0   # 1/20 pt
    hang = round(HANG * CONSOLAS_ADV * size * 20)
    keep = len(lines) <= 40
    last = len(lines) - 1
    for k, ln in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(2 if k == 0 else 0), Pt(8 if k == last else 0)
        pf.line_spacing = 1.08; pf.keep_together = True; pf.keep_with_next = keep and k < last
        ppr = p._p.get_or_add_pPr(); shade(ppr, "F4F3EE")
        para_border(p, ["top", "left", "bottom", "right"], "DCD8CC", sz=4, space=5)
        put(ppr, el("w:ind", left=110 + hang, right=110, hanging=hang))
        for chunk, wide in _wide_chunks(ln or " "):
            r = p.add_run(chunk)
            set_fonts(r._r, MONO, FONT); r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0x25, 0x2A, 0x27)
            if wide_pad and wide and "\uac00" <= chunk[0] <= "\ud7a3":
                r._r.get_or_add_rPr().append(el("w:spacing", val=wide_pad))

CONSOLAS_ADV, MALGUN_HANGUL_ADV = 0.5498, 1.0
HANG = 3                                            # 코드 줄이 접힐 때 이어지는 줄을 들여 쓰는 칸 수

def is_label(text):
    """**굵게** 만으로 된 짧은 문단 = 소제목 구실(부록 D 의 "제약"·"트리거" 등). 다음 내용과 같은 쪽에 둔다."""
    t = text.strip()
    return t.startswith("**") and t.endswith("**") and t.count("**") == 2 and dispw(t) <= 40

def is_kv(rows):
    """2열이고 첫 열이 짧은 표(항목/내용 형식) — 첫 열 폭을 고정해 표마다 들쭉날쭉하지 않게."""
    return all(len(r) == 2 for r in rows) and max(dispw(plain(r[0])) for r in rows) <= 14
KV_FIRST = Mm(32)

def _wide_chunks(line):
    """연속된 전각(W/F) 문자 덩어리와 나머지로 나눈다."""
    out = []
    for ch in line:
        wide = unicodedata.east_asian_width(ch) in "WF"
        if out and out[-1][1] == wide: out[-1][0] += ch
        else: out.append([ch, wide])
    return [(c, w) for c, w in out]

def add_quote(doc, paras):
    for k, lines in enumerate(paras):
        p = doc.add_paragraph()
        pf = p.paragraph_format; pf.space_after = Pt(6 if k == len(paras) - 1 else 2); pf.line_spacing = 1.25
        ppr = p._p.get_or_add_pPr(); shade(ppr, "F1EFE8")
        para_border(p, ["left"], "1A4B3A", sz=18, space=8)
        put(ppr, el("w:ind", left=180, right=60))
        for j, ln in enumerate(lines):
            if j: p.add_run().add_break()
            add_runs(p, ln, size=9.5)

def add_list(doc, items):
    for k, (marker, text) in enumerate(items):
        if marker in ("-", "*"):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, text)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Twips(400); p.paragraph_format.first_line_indent = Twips(-300)
            r = p.add_run(marker + " "); r.bold = True; r.font.color.rgb = PINE
            add_runs(p, text)
        p.paragraph_format.space_after = Pt(2 if k < len(items) - 1 else 6)

def heading(doc, level, text, bm=None):
    lvl = min(level - 1, 4)                         # ## → Heading 1
    p = doc.add_heading(level=lvl)
    if bm is not None:
        p._p.append(el("w:bookmarkStart", id=bm, name=f"_sl{bm}"))
    for t, b, it, code in inline(text):
        r = p.add_run(t)
        if code: set_fonts(r._r, MONO, FONT)
    if bm is not None:
        p._p.append(el("w:bookmarkEnd", id=bm))
    if lvl == 1 and text.startswith("부록"):
        p.paragraph_format.page_break_before = True
    return p

TOC_LEVELS = (2, 3)                                 # ## · ###

def add_toc(doc, entries):
    """TOC 필드 — 결과 칸에 항목·내부 링크를 미리 채운다. Word 에서 F9 를 누르면 Word 가 쪽번호까지 다시 만든다."""
    last = len(entries) - 1
    for k, (level, text) in enumerate(entries):
        p = doc.add_paragraph()
        pf = p.paragraph_format; pf.line_spacing = 1.15
        top = level == TOC_LEVELS[0]
        pf.space_before, pf.space_after = Pt(5 if top else 0), Pt(1)
        if not top: pf.left_indent = Twips(360)
        if k == 0:
            r = p.add_run(); r._r.append(el("w:fldChar", fldCharType="begin"))
            r = p.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
            it.text = ' TOC \\o "1-2" \\h \\z \\u '; r._r.append(it)
            r = p.add_run(); r._r.append(el("w:fldChar", fldCharType="separate"))
        h = el("w:hyperlink", anchor=f"_sl{k}", history=1)
        p._p.append(h)
        for t, b, it_, code in inline(text):
            r = p.add_run(t); r.bold = top or None; r.font.size = Pt(10 if top else 9)
            if code: set_fonts(r._r, MONO, FONT)
            h.append(r._r)
        if k == last:
            r = p.add_run(); r._r.append(el("w:fldChar", fldCharType="end"))

def build_docx(blocks, out_docx):
    doc = Document(); setup(doc)
    title = next(b[2] for b in blocks if b[0] == "h" and b[1] == 1)
    meta = next((b[1] for b in blocks if b[0] == "table"), None)
    version = next((r[1] for r in (meta or [])[1:] if r[0] == "문서 버전"), "")
    header_footer(doc, f"쉴더스랩 · 일반 서버 이관 명세서{(' v' + version) if version else ''}")
    cp = doc.core_properties
    cp.title, cp.author, cp.subject = plain(title), "쉴더스랩 (SHIELDUS LAB)", "일반 서버 이관 명세서"
    cp.keywords, cp.comments = "migration, PostgreSQL, nginx, Docker", "원본: deploy/SPEC.md — 이 파일을 직접 고치지 말 것"

    # 표지: 제목 → 메타 표 → 머리 인용 → 목차
    for _ in range(5): doc.add_paragraph()
    p = doc.add_paragraph(); r = p.add_run("SHIELDUS LAB"); r.font.size = Pt(10); r.bold = True
    r.font.color.rgb = PINE; r._r.get_or_add_rPr().append(el("w:spacing", val=40))
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(18); p.paragraph_format.line_spacing = 1.1
    head, _, tail = plain(title).partition(" — ")
    r = p.add_run(head); r.font.size = Pt(15); r.font.color.rgb = MUTED; r.bold = True
    if tail:
        r.add_break(); r = p.add_run(tail); r.font.size = Pt(26); r.bold = True; r.font.color.rgb = INK
    para_border(p, ["bottom"], "1A4B3A", sz=12, space=10)

    i = 0
    while i < len(blocks) and not (blocks[i][0] == "h" and blocks[i][1] == 1): i += 1
    i += 1
    while i < len(blocks) and blocks[i][0] in ("table", "quote"):
        b = blocks[i]
        if b[0] == "table": add_table(doc, b[1], widths=[int(KV_FIRST.twips), int(TEXT_W.twips - KV_FIRST.twips)])
        else: add_quote(doc, b[1])
        i += 1
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph(); r = p.add_run("목차"); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = PINE
    p.paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph(); r = p.add_run("제목을 누르면 해당 위치로 이동한다. 쪽번호는 Word 에서 목차를 선택하고 F9 → [전체 목차 업데이트].")
    r.font.size = Pt(8); r.font.color.rgb = MUTED; p.paragraph_format.space_after = Pt(8)
    entries = toc_entries(blocks)
    add_toc(doc, entries)
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)

    k = 0
    for b in blocks[i:]:
        kind = b[0]
        if kind == "h":
            if b[1] in TOC_LEVELS: heading(doc, b[1], b[2], bm=k); k += 1
            else: heading(doc, b[1], b[2])
        elif kind == "p":
            p = doc.add_paragraph(); add_runs(p, b[1])
            if is_label(b[1]): p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(2)
        elif kind == "table":
            add_table(doc, b[1], widths=[int(KV_FIRST.twips), int(TEXT_W.twips - KV_FIRST.twips)] if is_kv(b[1]) else None)
        elif kind == "code": add_code(doc, b[1])
        elif kind == "quote": add_quote(doc, b[1])
        elif kind == "list": add_list(doc, b[1])
        # hr: 장 제목이 바로 뒤따르므로 그리지 않는다
    doc.save(out_docx)
    return entries

def body_start(blocks):
    """제목(#) 다음의 표지 블록(메타 표·머리 인용)을 건너뛴 본문 시작 위치."""
    i = next(k for k, b in enumerate(blocks) if b[0] == "h" and b[1] == 1) + 1
    while i < len(blocks) and blocks[i][0] in ("table", "quote"): i += 1
    return i

def toc_entries(blocks):
    return [(b[1], b[2]) for b in blocks[body_start(blocks):] if b[0] == "h" and b[1] in TOC_LEVELS]


# ── PDF: 같은 블록 → HTML → Chrome 헤드리스 인쇄 ─────────────────────────────
def h_code(t):
    """인라인 코드 — 브라우저는 '-' 뒤에서 줄을 바꾼다(docker-|compose, --|zip). 명령·경로가 갈라져 보이므로
    공백 사이 조각을 줄바꿈 금지로 묶고, 경로 구분자 '/' 뒤에만 줄바꿈 기회(<wbr>)를 준다. 복사되는 글자는 그대로다."""
    out = []
    for piece in re.split(r"(\s+)", t):
        if not piece or piece.isspace(): out.append(html.escape(piece)); continue
        parts = re.split(r"(?<=/)(?=.)", piece)
        out.append("<wbr>".join(f'<span class="nb">{html.escape(x)}</span>' for x in parts))
    return "".join(out)

def h_text(t):
    """본문 — "(내부망)로만" 처럼 닫는 괄호 뒤에 조사가 붙은 어절은 브라우저가 ')' 뒤에서 끊는다
    (keep-all 이어도). 그런 어절만 줄바꿈 금지로 묶는다."""
    return re.sub(r"\S*\)[^\s)]+", lambda m: f'<span class="nb">{m.group(0)}</span>', html.escape(t))

def h_inline(text):
    out = []
    for t, b, it, code in inline(text):
        t = f"<code>{h_code(t)}</code>" if code else h_text(t)
        if it: t = f"<em>{t}</em>"
        if b: t = f"<strong>{t}</strong>"
        out.append(t)
    return "".join(out)

def is_diagram(lines):
    """상자 그림 문자(U+2500–257F)가 든 코드 블록 = 칸을 맞춰 그린 도식. 이것만 칸 고정을 한다
    (일반 코드의 한글 주석까지 2칸으로 벌리면 자간이 벌어져 보인다)."""
    return any("\u2500" <= ch <= "\u257f" for l in lines for ch in l)

def pre_line(line):
    """코드 블록 한 줄 — ASCII 밖 문자는 원문이 가정한 칸 수의 상자에 넣는다(글꼴 대체로 폭이 달라져도 도식이 맞게)."""
    out = []
    for ch in line:
        if ord(ch) < 128: out.append(html.escape(ch))
        else: out.append(f'<span class="w w{dispw(ch)}">{html.escape(ch)}</span>')
    return "".join(out)

CSS = """
@font-face { font-family: SLKR; src: url("__FONT__") format("truetype"); font-weight: 100 900; }
@page { size: A4; margin: 20mm 18mm 17mm 18mm;
  @top-right { content: "__LABEL__"; font: 600 7.5pt SLKR; color: #6B6A63; vertical-align: bottom;
               padding-bottom: 2.5mm; border-bottom: .5pt solid #D9D6CC; }
  @bottom-center { content: counter(page) " / " counter(pages); font: 7.5pt SLKR; color: #6B6A63; } }
@page :first { @top-right { content: none; border: none; } @bottom-center { content: none; } }
html { -webkit-print-color-adjust: exact; print-color-adjust: exact; }
body { font-family: SLKR, sans-serif; font-size: 9.4pt; line-height: 1.62; color: #1C1C1A; margin: 0;
       word-break: keep-all; overflow-wrap: break-word;
       font-variant-ligatures: none; font-feature-settings: "liga" 0, "clig" 0; }  /* PDF 에서 복사한 명령이 ﬀ·ﬁ 로 바뀌지 않게 */
strong { font-weight: 700; } em { font-style: italic; }
code { font-family: Menlo, Consolas, "SF Mono", SLKR, monospace; font-size: .86em; background: #EFEDE6;
       padding: .05em .28em; border-radius: 2px; word-break: normal; overflow-wrap: break-word; }
h2, h3, h4 { break-after: avoid; line-height: 1.3; }
h2 { font-size: 15pt; font-weight: 700; color: #1A4B3A; margin: 20pt 0 7pt; padding-bottom: 4pt;
     border-bottom: 1pt solid #1A4B3A; }
h3 { font-size: 11.6pt; font-weight: 700; margin: 14pt 0 5pt; }
h4 { font-size: 10.2pt; font-weight: 700; margin: 10pt 0 4pt; }
h2 code, h3 code { font-size: .82em; }
h2.appendix { break-before: page; }
p { margin: 0 0 5pt; }
ul, ol { margin: 0 0 7pt; padding-left: 16pt; } li { margin: 0 0 2pt; } ol { list-style: none; padding-left: 4pt; }
ol li .n { color: #1A4B3A; font-weight: 700; margin-right: 4pt; }
table { width: 100%; border-collapse: collapse; font-size: 7.9pt; line-height: 1.42; margin: 3pt 0 10pt; }
thead { display: table-header-group; } tr { break-inside: avoid; }
th, td { border: .5pt solid #CFCBC0; padding: 2.6pt 4.5pt; text-align: left; vertical-align: top; }
th { background: #1A4B3A; color: #fff; font-weight: 700; }
td code { font-size: .92em; }
pre { background: #F4F3EE; border: .5pt solid #DCD8CC; border-radius: 2px; padding: 6pt 8pt; margin: 2pt 0 9pt;
      font-family: Menlo, Consolas, "SF Mono", SLKR, monospace; font-size: 7.2pt; line-height: 1.45;
      white-space: pre-wrap; overflow-wrap: anywhere; color: #252A27; }
pre .ln { display: block; padding-left: __HANG__ch; text-indent: -__HANG__ch; }  /* 접힌 줄은 들여 써서 새 줄과 구분 */
code .nb, .nb { white-space: nowrap; }
p.label { break-after: avoid; margin: 8pt 0 2pt; }
table.kv td:first-child, table.kv th:first-child { width: 32mm; }
pre .w { display: inline-block; text-align: center; text-indent: 0; }  /* 비ASCII 문자를 원문 칸 수(한글 2칸·기호 1칸)에 고정 · .ln 의 음수 들여쓰기 상속 차단 */
pre .w1 { width: 1ch; } pre .w2 { width: 2ch; }
td.nw, th.nw { white-space: nowrap; }
pre.short { break-inside: avoid; }
blockquote { background: #F1EFE8; border-left: 2.6pt solid #1A4B3A; margin: 3pt 0 10pt; padding: 6pt 10pt;
             font-size: 9pt; line-height: 1.55; break-inside: avoid; }
blockquote p { margin: 0 0 3pt; } blockquote p:last-child { margin: 0; }
.cover { height: 247mm; display: flex; flex-direction: column; justify-content: center; break-after: page; }
.cover .brand { font-size: 9.5pt; font-weight: 700; letter-spacing: .32em; color: #1A4B3A; margin-bottom: 8pt; }
.cover .t1 { font-size: 15pt; font-weight: 700; color: #6B6A63; }
.cover .t2 { font-size: 27pt; font-weight: 700; letter-spacing: -.01em; margin: 2pt 0 16pt; padding-bottom: 12pt;
             border-bottom: 1.6pt solid #1A4B3A; }
.toc { break-after: page; }
.toc .toc-title { font-size: 16pt; font-weight: 700; color: #1A4B3A; margin: 0 0 10pt; }
.toc a { display: flex; align-items: baseline; color: inherit; text-decoration: none; }
.toc .l2 a { font-weight: 700; font-size: 9.8pt; margin-top: 5pt; }
.toc .l3 a { font-size: 8.8pt; padding-left: 14pt; color: #33332F; }
.toc .dots { flex: 1; border-bottom: .8pt dotted #A8A496; margin: 0 4pt 2.6pt; min-width: 8pt; }
.toc .pg { flex: none; min-width: 12pt; text-align: right; }
"""

def build_html(blocks, pages=None):
    title = next(b[2] for b in blocks if b[0] == "h" and b[1] == 1)
    meta = next((b[1] for b in blocks if b[0] == "table"), None)
    version = next((r[1] for r in (meta or [])[1:] if r[0] == "문서 버전"), "")
    label = f"쉴더스랩 · 일반 서버 이관 명세서{(' v' + version) if version else ''}"
    font = pathlib.Path(os.path.dirname(os.path.dirname(HERE)), "tools", "NotoSansKR.ttf").as_uri()
    css = CSS.replace("__FONT__", font).replace("__LABEL__", label).replace("__HANG__", str(HANG))

    def cell(tag, c):
        nw = ' class="nw"' if 0 < dispw(plain(c)) <= 12 else ""
        return f"<{tag}{nw}>{h_inline(c)}</{tag}>"

    def table(rows, cls=""):
        head, body = rows[0], rows[1:]
        h = "".join(cell("th", c) for c in head)
        b = "".join("<tr>" + "".join(cell("td", c) for c in r) + "</tr>" for r in body)
        return f'<table class="{cls}"><thead><tr>{h}</tr></thead><tbody>{b}</tbody></table>'

    def block(b):
        kind = b[0]
        if kind == "p": return f'<p{" class=label" if is_label(b[1]) else ""}>{h_inline(b[1])}</p>'
        if kind == "table": return table(b[1], "kv" if is_kv(b[1]) else "")
        if kind == "code":
            fix = is_diagram(b[1])
            body = "".join('<span class="ln">' + ((pre_line(l) if fix else html.escape(l)) or " ") + "</span>" for l in b[1])
            return f'<pre class="{"short" if len(b[1]) <= 40 else ""}">' + body + "</pre>"
        if kind == "quote":
            return "<blockquote>" + "".join("<p>" + "<br>".join(h_inline(l) for l in para) + "</p>" for para in b[1]) + "</blockquote>"
        if kind == "list":
            if all(m in ("-", "*") for m, _ in b[1]):
                return "<ul>" + "".join(f"<li>{h_inline(t)}</li>" for _, t in b[1]) + "</ul>"
            return "<ol>" + "".join(f'<li><span class="n">{html.escape(m)}</span>{h_inline(t)}</li>' for m, t in b[1]) + "</ol>"
        return ""

    out = [f'<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>{html.escape(plain(title))}</title>'
           f"<style>{css}</style></head><body>"]
    head, _, tail = plain(title).partition(" — ")
    start = body_start(blocks)
    i = next(k for k, b in enumerate(blocks) if b[0] == "h" and b[1] == 1) + 1
    out.append(f'<section class="cover"><div class="brand">SHIELDUS LAB</div><div class="t1">{html.escape(head)}</div>'
               f'<div class="t2">{html.escape(tail or head)}</div>'
               + "".join(table(b[1], "kv") if b[0] == "table" else block(b) for b in blocks[i:start]) + "</section>")
    entries = toc_entries(blocks)
    toc = ['<section class="toc"><div class="toc-title">목차</div>']   # h1 을 쓰면 PDF 책갈피가 전부 "목차" 밑으로 들어간다
    for k, (lvl, text) in enumerate(entries):
        pg = pages[k] if pages and k < len(pages) and pages[k] else ""
        toc.append(f'<div class="l{lvl}"><a href="#s{k}"><span>{h_inline(text)}</span>'
                   f'<span class="dots"></span><span class="pg">{pg}</span></a></div>')
    out.append("".join(toc) + "</section>")
    k = 0
    for b in blocks[start:]:
        if b[0] == "h":
            lvl = min(b[1], 4)
            attrs = ""
            if b[1] in TOC_LEVELS: attrs = f' id="s{k}"'; k += 1
            if lvl == 2 and b[2].startswith("부록"): attrs += ' class="appendix"'
            out.append(f"<h{lvl}{attrs}>{h_inline(b[2])}</h{lvl}>")
        else:
            out.append(block(b))
    out.append("</body></html>")
    return "\n".join(out), entries

def find_chrome():
    cands = [os.environ.get("SL_CHROME", ""),
             "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
             "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
             r"C:\Program Files\Google\Chrome\Application\chrome.exe",
             r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
             r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
             r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"]
    cands += [shutil.which(n) or "" for n in ("google-chrome", "chromium", "chromium-browser", "msedge")]
    return next((c for c in cands if c and os.path.exists(c)), None)

def print_pdf(chrome, html_text, out_pdf, work):
    src = os.path.join(work, "spec.html")
    open(src, "w", encoding="utf-8").write(html_text)
    if os.path.exists(out_pdf): os.remove(out_pdf)
    cmd = [chrome, "--headless", "--disable-gpu", "--no-first-run", "--no-default-browser-check",
           f"--user-data-dir={os.path.join(work, 'profile')}", "--allow-file-access-from-files",
           "--no-pdf-header-footer", "--generate-pdf-document-outline",
           f"--print-to-pdf={out_pdf}", pathlib.Path(src).as_uri()]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    if not os.path.exists(out_pdf):
        raise RuntimeError("Chrome 인쇄 실패: " + (r.stderr or r.stdout).strip()[-300:])

def find_pages(pdf, entries):
    """PDF 에서 각 목차 항목이 처음 나오는 쪽(1부터). 목차 쪽 자체는 건너뛴다."""
    from pypdf import PdfReader
    norm = lambda t: re.sub(r"\s+", "", t)
    texts = [norm(pg.extract_text() or "") for pg in PdfReader(pdf).pages]
    first = norm(plain(entries[0][1]))
    hits = [n for n, t in enumerate(texts) if first in t]
    cur = hits[1] if len(hits) > 1 else (hits[0] if hits else 0)   # 첫 번째 일치는 목차 쪽
    pages = []
    for _, text in entries:
        key = norm(plain(text))
        n = next((j for j in range(cur, len(texts)) if key in texts[j]), None)
        if n is None: pages.append(None); continue
        pages.append(n + 1); cur = n
    return pages, len(texts)

def build_pdf(blocks, out_pdf):
    chrome = find_chrome()
    if not chrome: raise RuntimeError("Chrome/Edge 를 못 찾았다(SL_CHROME 로 경로 지정)")
    with tempfile.TemporaryDirectory(prefix="sl-spec-") as work:
        pages, rnd = None, 0
        for rnd in range(1, 5):
            text, entries = build_html(blocks, pages)
            print_pdf(chrome, text, out_pdf, work)
            found, npages = find_pages(out_pdf, entries)
            if found == pages: break                # 쪽번호를 넣어도 배치가 그대로 → 수렴
            pages = found
        else:
            print("  ⚠ 목차 쪽번호가 4회 안에 수렴하지 않았다 — 확인할 것")
        if os.environ.get("SL_KEEP_HTML"): shutil.copy(os.path.join(work, "spec.html"), out_pdf[:-4] + ".html")
    missing = [entries[k][1] for k, p in enumerate(pages or []) if not p]
    return npages, rnd, missing

# ── 검증: 원문 조각이 결과물에 빠짐없이 들어갔는가 ─────────────────────────
def source_items(blocks):
    """원문을 비교 단위로 쪼갠다 — 문단·목록 항목·인용 줄·표 칸·코드 줄·제목."""
    items = []
    for b in blocks:
        if b[0] == "h" and b[1] == 1:                                    # 표지에서는 " — " 로 나눠 두 줄로 쓴다
            items += [("표제", t) for t in plain(b[2]).split(" — ")]
        elif b[0] == "h": items.append(("제목", plain(b[2])))
        elif b[0] == "p": items.append(("문단", plain(b[1])))
        elif b[0] == "list": items += [("목록", plain(t)) for _, t in b[1]]
        elif b[0] == "quote": items += [("인용", plain(l)) for para in b[1] for l in para]
        elif b[0] == "table": items += [("표", plain(c)) for r in b[1] for c in r if c.strip()]
        elif b[0] == "code": items += [("코드", l) for l in b[1] if l.strip()]
    return items

def _squash(t):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", t))

def verify(blocks, docx_path, pdf_path=None):
    """원문 조각마다 docx·pdf 텍스트에 들어 있는지 본다. 마크다운 기호가 새어 나왔는지도 본다."""
    items = source_items(blocks)
    d = Document(docx_path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells: parts.append(c.text)
    dtext = _squash("".join(parts))
    report = {"docx": [k + ": " + v for k, v in items if _squash(v) not in dtext]}
    leak = [m for m in ("**", "\\|") if m in "".join(p.text for p in d.paragraphs if not p.text.startswith(" "))]
    if leak: report["docx"].append(f"마크다운 기호 노출: {leak}")
    if pdf_path:
        from pypdf import PdfReader
        pages = []
        for pg in PdfReader(pdf_path).pages:
            t = pg.extract_text() or ""
            t = re.sub(r"쉴더스랩 · 일반 서버 이관 명세서 v[\d.]+", "", t)
            t = re.sub(r"(?m)^\s*\d+ / \d+\s*$", "", t)
            pages.append(t)
        ptext = _squash("".join(pages))
        report["pdf"] = [k + ": " + v for k, v in items if _squash(v) not in ptext]
    return len(items), report

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("outdir"); ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--spec", default=SPEC)
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    out_docx = os.path.join(a.outdir, BASENAME + ".docx")
    out_pdf = os.path.join(a.outdir, BASENAME + ".pdf")
    blocks = parse(open(a.spec, encoding="utf-8").read())
    kinds = {}
    for b in blocks: kinds[b[0]] = kinds.get(b[0], 0) + 1
    entries = build_docx(blocks, out_docx)
    print("  · 블록 " + " · ".join(f"{k} {v}" for k, v in sorted(kinds.items())) + f" · 목차 {len(entries)}항목")
    print(f"  ✓ {os.path.basename(out_docx)}")
    if a.no_pdf: return 0
    try:
        npages, rnd, missing = build_pdf(blocks, out_pdf)
    except (RuntimeError, ImportError, subprocess.TimeoutExpired) as e:
        print(f"  ⚠ PDF 건너뜀 — {e}"); return 2
    print(f"  ✓ {os.path.basename(out_pdf)} — {npages}쪽 · 목차 쪽번호 {rnd}회차에 확정")
    if missing:
        print(f"  ✗ PDF 에서 못 찾은 제목 {len(missing)}개: {missing[:3]}"); return 1
    n, rep = verify(blocks, out_docx, out_pdf)
    bad = {k: v for k, v in rep.items() if v}
    if bad:
        for k, v in bad.items():
            print(f"  ✗ {k} 누락 {len(v)}건: " + " | ".join(x[:60] for x in v[:8]))
        return 1
    print(f"  ✓ 원문 대조 — 조각 {n}개 전부 docx·pdf 에 있음")
    return 0

if __name__ == "__main__":
    sys.exit(main())
